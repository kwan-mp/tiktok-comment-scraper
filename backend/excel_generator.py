import io

from deep_translator import GoogleTranslator
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["序号", "用户名", "创作者", "评论内容", "评论内容（中文）", "发布时间", "点赞数", "用户主页"]
COLUMN_WIDTHS = [6, 20, 8, 40, 40, 12, 8, 32]


def translate_to_chinese(text):
    if not text or not text.strip():
        return ''
    try:
        return GoogleTranslator(source='auto', target='zh-CN').translate(text)
    except Exception:
        return ''


def build_excel(comments):
    wb = Workbook()
    ws = wb.active
    ws.title = "评论明细"

    ws.append(HEADERS)
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_idx in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    body_font = Font(name="Arial", size=11)

    for i, c in enumerate(comments, start=1):
        comment_text = c.get('comment', '')
        username = c.get('username', '')
        row = [
            i,
            username,
            '是' if c.get('isCreator') else '否',
            comment_text,
            translate_to_chinese(comment_text),
            (c.get('date') or '')[:10],
            c.get('likes', 0),
            f'https://www.tiktok.com/{username}' if username else '',
        ]
        ws.append(row)
        for col_idx in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=i + 1, column=col_idx)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=(col_idx in (4, 5)))

    for idx, w in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
